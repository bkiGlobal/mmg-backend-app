import json
from io import BytesIO
from datetime import date, datetime, time
from tempfile import TemporaryDirectory

from django.contrib import admin
from django.contrib.auth.models import Permission, User
from django.contrib.gis.geos import Point
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import RequestFactory, TestCase, override_settings
from django.templatetags.static import static
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image
from rest_framework.test import APIClient

from core.models import Location
from project.models import Project

from .models import (
    Attendance,
    AttendanceStatus,
    AttendanceWorkMode,
    GenderType,
    Holiday,
    HolidaySource,
    Notifications,
    Profile,
    RoleType,
    StatusType,
    Team,
    TeamMember,
    WorkPolicy,
)
from .holidays import fetch_indonesian_holidays, sync_indonesian_holidays
from .services import generate_daily_attendance


class SoftDeleteTests(TestCase):
    def setUp(self):
        self.actor = User.objects.create_user(
            username='soft-delete-actor',
            password='test-password',
        )
        self.profile = Profile.objects.create(
            user=self.actor,
            full_name='Soft Delete Actor',
            role=RoleType.ADMIN,
            gender=GenderType.MALE,
            status=StatusType.PERMANENT,
            birthday=date(1990, 1, 1),
            join_date=date(2024, 1, 1),
            phone_number='08123456789',
        )
        self.team = Team.objects.create(
            name='Test Team',
            description='Soft delete test team',
        )
        self.member = TeamMember.objects.create(
            team=self.team,
            user=self.profile,
        )

    def test_model_delete_hides_row_and_cascades_to_audit_children(self):
        self.team.delete(user=self.actor)

        self.assertFalse(Team.objects.filter(pk=self.team.pk).exists())
        self.assertTrue(
            Team.all_objects.get(pk=self.team.pk).is_deleted
        )
        self.assertFalse(
            TeamMember.objects.filter(pk=self.member.pk).exists()
        )
        self.assertTrue(
            TeamMember.all_objects.get(pk=self.member.pk).is_deleted
        )

    def test_queryset_delete_is_soft_delete(self):
        Team.objects.filter(pk=self.team.pk).delete(user=self.actor)

        self.assertTrue(
            Team.all_objects.filter(pk=self.team.pk).exists()
        )
        self.assertFalse(Team.objects.filter(pk=self.team.pk).exists())

    def test_restore_recovers_parent_and_cascade_children(self):
        self.team.delete(user=self.actor)
        deleted_team = Team.all_objects.get(pk=self.team.pk)

        deleted_team.restore(user=self.actor)

        self.assertTrue(Team.objects.filter(pk=self.team.pk).exists())
        self.assertTrue(
            TeamMember.objects.filter(pk=self.member.pk).exists()
        )

    def test_restore_does_not_revive_previously_deleted_child(self):
        self.member.delete(user=self.actor)
        self.team.delete(user=self.actor)

        Team.all_objects.get(pk=self.team.pk).restore(user=self.actor)

        self.assertTrue(Team.objects.filter(pk=self.team.pk).exists())
        self.assertFalse(
            TeamMember.objects.filter(pk=self.member.pk).exists()
        )

    def test_admin_bulk_delete_and_restore_ignore_nullable_outer_joins(self):
        attendance = Attendance.objects.create(
            user=self.profile,
            date=date(2026, 7, 28),
            status=AttendanceStatus.ABSENT,
        )
        model_admin = admin.site._registry[Attendance]
        request = RequestFactory().post('/admin/team/attendance/')
        request.user = self.actor
        queryset = model_admin.get_queryset(request).filter(
            pk=attendance.pk,
        )

        self.assertIn('LEFT OUTER JOIN', str(queryset.query))
        model_admin.delete_queryset(request, queryset)

        attendance.refresh_from_db()
        self.assertTrue(attendance.is_deleted)
        self.assertEqual(attendance.deleted_by, self.actor)

        restore_request = RequestFactory().get(
            '/admin/team/attendance/',
            {'is_deleted__exact': '1'},
        )
        restore_request.user = self.actor
        restore_queryset = model_admin.get_queryset(
            restore_request
        ).filter(pk=attendance.pk)

        self.assertEqual(
            model_admin.restore_queryset(
                restore_request,
                restore_queryset,
            ),
            1,
        )
        attendance.refresh_from_db()
        self.assertFalse(attendance.is_deleted)
        self.assertIsNone(attendance.deleted_by)

    def test_soft_delete_fix_applies_to_nullable_joins_in_other_apps(self):
        project = Project.objects.create(
            project_name='Cross App Delete',
            project_code='DELETE-001',
            description='Regression test for nullable admin joins.',
            start_date=date(2026, 7, 28),
        )
        joined_queryset = Project.all_objects.select_related(
            'location',
            'client',
            'team',
        ).filter(pk=project.pk)

        self.assertIn('LEFT OUTER JOIN', str(joined_queryset.query))
        deleted_count, deleted_per_model = joined_queryset.delete(
            user=self.actor,
        )

        self.assertEqual(deleted_count, 1)
        self.assertEqual(deleted_per_model, {'project.Project': 1})
        self.assertTrue(
            Project.all_objects.get(pk=project.pk).is_deleted
        )

        restored_count = (
            Project.all_objects
            .select_related('location', 'client', 'team')
            .filter(pk=project.pk)
            .restore(user=self.actor)
        )

        self.assertEqual(restored_count, 1)
        self.assertTrue(Project.objects.filter(pk=project.pk).exists())


class ProfileVisibilityAndOwnershipTests(TestCase):
    def setUp(self):
        self.staff_user = User.objects.create_user(
            username='profile-owner',
            password='test-password',
            is_staff=True,
        )
        self.staff_profile = Profile.objects.create(
            user=self.staff_user,
            full_name='Profile Owner',
            role=RoleType.WORKER,
            gender=GenderType.MALE,
            status=StatusType.PERMANENT,
            birthday=date(1990, 1, 1),
            join_date=date(2024, 1, 1),
            phone_number='0811111111',
        )
        self.other_user = User.objects.create_user(
            username='other-profile',
            password='test-password',
            is_staff=True,
        )
        self.other_profile = Profile.objects.create(
            user=self.other_user,
            full_name='Other Staff',
            role=RoleType.ARCHITECT,
            gender=GenderType.FEMALE,
            status=StatusType.PERMANENT,
            birthday=date(1992, 2, 2),
            join_date=date(2024, 2, 2),
            phone_number='0822222222',
        )
        self.profile_admin = admin.site._registry[Profile]

    def test_admin_staff_can_view_all_but_only_change_own_profile(self):
        request = RequestFactory().get('/admin/team/profile/')
        request.user = self.staff_user

        visible_profiles = self.profile_admin.get_queryset(request)
        self.assertSetEqual(
            set(visible_profiles.values_list('pk', flat=True)),
            {self.staff_profile.pk, self.other_profile.pk},
        )
        self.assertTrue(
            self.profile_admin.has_view_permission(
                request,
                self.other_profile,
            )
        )
        self.assertTrue(
            self.profile_admin.has_change_permission(
                request,
                self.staff_profile,
            )
        )
        self.assertFalse(
            self.profile_admin.has_change_permission(
                request,
                self.other_profile,
            )
        )
        self.assertFalse(
            self.profile_admin.has_delete_permission(
                request,
                self.staff_profile,
            )
        )
        self.assertIn(
            'role',
            self.profile_admin.get_readonly_fields(
                request,
                self.staff_profile,
            ),
        )

        client = APIClient()
        client.force_login(self.staff_user)
        list_response = client.get(
            '/admin/team/profile/',
            secure=True,
        )
        other_response = client.get(
            f'/admin/team/profile/{self.other_profile.pk}/change/',
            secure=True,
        )
        denied_change = client.post(
            f'/admin/team/profile/{self.other_profile.pk}/change/',
            {'full_name': 'Tidak Boleh Berubah'},
            secure=True,
        )

        self.assertEqual(list_response.status_code, 200)
        self.assertContains(list_response, self.staff_profile.full_name)
        self.assertContains(list_response, self.other_profile.full_name)
        self.assertEqual(other_response.status_code, 200)
        self.assertContains(other_response, 'Profil ini hanya dapat dilihat')
        self.assertNotContains(other_response, 'id_birthday')
        self.assertNotContains(other_response, 'id_location')
        self.assertEqual(denied_change.status_code, 403)
        self.other_profile.refresh_from_db()
        self.assertEqual(self.other_profile.full_name, 'Other Staff')

    def test_profile_api_is_directory_and_update_is_owner_only(self):
        client = APIClient()
        client.force_authenticate(self.staff_user)

        list_response = client.get('/api/team/profile/', secure=True)
        self.assertEqual(list_response.status_code, 200)
        listed_ids = {
            item['id'] for item in list_response.data['results']
        }
        self.assertSetEqual(
            listed_ids,
            {
                str(self.staff_profile.pk),
                str(self.other_profile.pk),
            },
        )

        other_response = client.get(
            f'/api/team/profile/{self.other_profile.pk}/',
            secure=True,
        )
        self.assertEqual(other_response.status_code, 200)
        self.assertNotIn('birthday', other_response.data)
        self.assertNotIn('location', other_response.data)
        self.assertNotIn('user_attendance', other_response.data)

        update_payload = {
            'full_name': 'Nama Pribadi Baru',
            'gender': GenderType.MALE,
            'birthday': '1990-01-01',
            'phone_number': '0833333333',
            'role': RoleType.ADMIN,
        }
        denied_response = client.put(
            f'/api/team/profile/{self.other_profile.pk}/',
            update_payload,
            format='json',
            secure=True,
        )
        self.assertEqual(denied_response.status_code, 403)
        denied_create = client.post(
            '/api/team/profile/',
            {},
            format='json',
            secure=True,
        )
        self.assertEqual(denied_create.status_code, 403)

        own_response = client.put(
            f'/api/team/profile/{self.staff_profile.pk}/',
            update_payload,
            format='json',
            secure=True,
        )
        self.assertEqual(own_response.status_code, 200)
        self.staff_profile.refresh_from_db()
        self.other_profile.refresh_from_db()
        self.assertEqual(
            self.staff_profile.full_name,
            'Nama Pribadi Baru',
        )
        self.assertEqual(
            self.staff_profile.phone_number,
            '0833333333',
        )
        self.assertEqual(self.staff_profile.role, RoleType.WORKER)
        self.assertEqual(self.other_profile.full_name, 'Other Staff')


class AttendanceAutomationTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='attendance-user',
            password='test-password',
        )
        self.policy = WorkPolicy.objects.create(
            name='Attendance Test',
            work_start=time(9, 0),
            work_end=time(17, 0),
            late_grace_minutes=10,
            overtime_after_minutes=30,
            workdays=[0, 1, 2, 3, 4],
        )
        self.profile = Profile.objects.create(
            user=self.user,
            full_name='Attendance User',
            role=RoleType.WORKER,
            gender=GenderType.MALE,
            status=StatusType.PERMANENT,
            birthday=date(1990, 1, 1),
            join_date=date(2024, 1, 1),
            phone_number='0811111111',
            work_policy=self.policy,
        )

    def test_policy_calculates_late_worked_and_overtime_minutes(self):
        attendance_date = date(2026, 8, 3)
        check_in = timezone.make_aware(
            datetime.combine(attendance_date, time(9, 15))
        )
        check_out = timezone.make_aware(
            datetime.combine(attendance_date, time(18, 0))
        )
        attendance = Attendance.objects.create(
            user=self.profile,
            date=attendance_date,
            check_in=check_in,
            check_out=check_out,
            status=AttendanceStatus.ONTIME,
            work_policy=self.policy,
        )
        self.assertEqual(attendance.status, AttendanceStatus.LATE)
        self.assertEqual(attendance.worked_minutes, 525)
        self.assertEqual(attendance.overtime_minutes, 60)

    def test_daily_automation_creates_absence_once_and_notifies(self):
        monday = date(2026, 8, 3)
        self.assertEqual(generate_daily_attendance(monday), 1)
        self.assertEqual(generate_daily_attendance(monday), 0)
        attendance = Attendance.objects.get(
            user=self.profile,
            date=monday,
        )
        self.assertEqual(attendance.status, AttendanceStatus.ABSENT)
        self.assertTrue(
            Notifications.objects.filter(
                user=self.profile,
                category='attendance',
            ).exists()
        )

    def test_superuser_status_override_wins_over_calculation(self):
        attendance_date = date(2026, 8, 3)
        attendance = Attendance.objects.create(
            user=self.profile,
            date=attendance_date,
            check_in=timezone.make_aware(
                datetime.combine(attendance_date, time(9, 0))
            ),
            status=AttendanceStatus.ONTIME,
            status_override=AttendanceStatus.LEAVE,
            status_override_reason='Koreksi berdasarkan surat cuti.',
            work_policy=self.policy,
        )

        self.assertEqual(attendance.status, AttendanceStatus.LEAVE)


class HolidayCalendarTests(TestCase):
    """Sinkronisasi kalender libur dan interaksinya dengan absensi."""

    def setUp(self):
        self.user = User.objects.create_user(
            username='holiday-user',
            password='holiday-pass',
        )
        self.policy = WorkPolicy.objects.create(
            name='Holiday Policy',
            work_start=time(9, 0),
            work_end=time(17, 0),
            late_grace_minutes=10,
            overtime_after_minutes=30,
            workdays=[0, 1, 2, 3, 4],
        )
        self.profile = Profile.objects.create(
            user=self.user,
            full_name='Holiday User',
            role=RoleType.WORKER,
            gender=GenderType.MALE,
            status=StatusType.PERMANENT,
            birthday=date(1990, 1, 1),
            join_date=date(2024, 1, 1),
            phone_number='0812222222',
            work_policy=self.policy,
        )

    def test_source_covers_lunar_and_hijri_holidays(self):
        """Kalender nasional harus memuat libur non-Masehi."""
        entries = dict(fetch_indonesian_holidays(2026))
        self.assertIn(date(2026, 3, 19), entries)   # Nyepi
        self.assertIn(date(2026, 5, 27), entries)   # Idul Adha
        self.assertIn(date(2026, 8, 17), entries)   # Kemerdekaan
        self.assertGreaterEqual(len(entries), 15)

    def test_sync_creates_national_entries_then_is_idempotent(self):
        summary = sync_indonesian_holidays(2026, 8)
        self.assertTrue(summary['created'])
        holiday = Holiday.objects.get(date=date(2026, 8, 17))
        self.assertEqual(holiday.source, HolidaySource.NATIONAL)

        repeat = sync_indonesian_holidays(2026, 8)
        self.assertEqual(repeat['created'], [])
        self.assertEqual(repeat['updated'], [])
        self.assertTrue(repeat['unchanged'])

    def test_sync_never_overwrites_manual_entry(self):
        sync_indonesian_holidays(2026, 8)
        holiday = Holiday.objects.get(date=date(2026, 8, 17))
        holiday.source = HolidaySource.MANUAL
        holiday.name = 'Libur perusahaan'
        holiday.save()

        summary = sync_indonesian_holidays(2026, 8)

        self.assertIn(
            (date(2026, 8, 17), 'Libur perusahaan'),
            summary['skipped_manual'],
        )
        holiday.refresh_from_db()
        self.assertEqual(holiday.name, 'Libur perusahaan')

    def test_sync_does_not_resurrect_deleted_entry(self):
        sync_indonesian_holidays(2026, 8)
        holiday = Holiday.objects.get(date=date(2026, 8, 17))
        holiday.is_deleted = True
        holiday.save()

        summary = sync_indonesian_holidays(2026, 8)

        self.assertTrue(summary['skipped_deleted'])
        self.assertFalse(
            Holiday.objects.filter(date=date(2026, 8, 17)).exists()
        )

    def test_workday_holiday_creates_holiday_attendance(self):
        monday = date(2026, 8, 3)
        Holiday.objects.create(date=monday, name='Libur uji')

        self.assertEqual(generate_daily_attendance(monday), 1)

        attendance = Attendance.objects.get(user=self.profile, date=monday)
        self.assertEqual(attendance.status, AttendanceStatus.HOLYDAY)

    def test_weekend_holiday_creates_no_attendance(self):
        saturday = date(2026, 8, 8)
        self.assertEqual(saturday.weekday(), 5)
        Holiday.objects.create(date=saturday, name='Libur akhir pekan')

        self.assertEqual(generate_daily_attendance(saturday), 0)
        self.assertFalse(
            Attendance.objects.filter(
                user=self.profile,
                date=saturday,
            ).exists()
        )

    def test_working_on_holiday_is_counted_fully_as_overtime(self):
        holiday_date = date(2026, 8, 17)
        Holiday.objects.create(date=holiday_date, name='Hari Kemerdekaan')
        check_in = timezone.make_aware(
            datetime.combine(holiday_date, time(9, 0))
        )
        check_out = timezone.make_aware(
            datetime.combine(holiday_date, time(13, 0))
        )

        attendance = Attendance.objects.create(
            user=self.profile,
            date=holiday_date,
            check_in=check_in,
            check_out=check_out,
            status=AttendanceStatus.ONTIME,
            work_policy=self.policy,
        )

        self.assertEqual(attendance.status, AttendanceStatus.HOLYDAY)
        self.assertEqual(attendance.worked_minutes, 240)
        self.assertEqual(attendance.overtime_minutes, 240)


class HolidayAdminCalendarTests(TestCase):
    """Kalender pada changelist Holiday beserta tombol sinkronisasinya."""

    def setUp(self):
        self.superuser = User.objects.create_superuser(
            username='holiday-admin',
            email='holiday@example.com',
            password='holiday-pass',
        )
        self.client.force_login(self.superuser)
        self.url = reverse('admin:team_holiday_changelist')

    def test_calendar_navigation_parameters_do_not_break_changelist(self):
        """cal_year/cal_month bukan filter, jadi harus dibuang sebelum
        changelist memvalidasi querystring."""
        response = self.client.get(
            self.url,
            {'cal_year': '2026', 'cal_month': '8'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Agustus 2026')

    def test_calendar_year_is_rendered_without_thousand_separator(self):
        response = self.client.get(
            self.url,
            {'cal_year': '2026', 'cal_month': '8'},
        )

        self.assertContains(response, 'value="2026"')
        self.assertNotContains(response, 'value="2.026"')

    def test_sync_button_imports_national_holidays(self):
        response = self.client.post(
            reverse('admin:team_holiday_sync'),
            {'year': '2026', 'month': '8', 'scope': 'month'},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        holiday = Holiday.objects.get(date=date(2026, 8, 17))
        self.assertEqual(holiday.source, HolidaySource.NATIONAL)

    def test_sync_rejects_get_request(self):
        response = self.client.get(reverse('admin:team_holiday_sync'))

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Holiday.objects.exists())


class AttendanceHeatmapVisibilityTests(TestCase):
    """Heatmap absensi hanya memuat staff, bukan akun client."""

    def setUp(self):
        self.superuser = User.objects.create_superuser(
            username='heatmap-admin',
            email='heatmap@example.com',
            password='heatmap-pass',
        )
        self.client.force_login(self.superuser)
        self.policy = WorkPolicy.objects.create(
            name='Heatmap Policy',
            work_start=time(9, 0),
            work_end=time(17, 0),
            workdays=[0, 1, 2, 3, 4],
        )
        self.staff = self._profile('Staff Lapangan', RoleType.WORKER, '1')
        self.client_profile = self._profile(
            'Klien Proyek',
            RoleType.CLIENT,
            '2',
        )

    def _profile(self, full_name, role, suffix):
        user = User.objects.create_user(
            username=f'heatmap-user-{suffix}',
            password='heatmap-pass',
        )
        return Profile.objects.create(
            user=user,
            full_name=full_name,
            role=role,
            gender=GenderType.MALE,
            status=StatusType.PERMANENT,
            birthday=date(1990, 1, 1),
            join_date=date(2024, 1, 1),
            phone_number=f'0813333{suffix}',
            work_policy=self.policy,
        )

    def test_heatmap_lists_staff_but_not_client(self):
        response = self.client.get(
            reverse('admin:team_attendance_changelist')
        )

        self.assertEqual(response.status_code, 200)
        heatmap = response.context_data['attendance_heatmap']
        names = [row['profile'].full_name for row in heatmap['rows']]
        self.assertIn('Staff Lapangan', names)
        self.assertNotIn('Klien Proyek', names)

    def test_heatmap_hidden_for_client_account(self):
        self.client.force_login(self.client_profile.user)
        permission = Permission.objects.get(codename='view_attendance')
        self.client_profile.user.is_staff = True
        self.client_profile.user.save()
        self.client_profile.user.user_permissions.add(permission)

        response = self.client.get(
            reverse('admin:team_attendance_changelist')
        )

        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.context_data['attendance_heatmap'])


class AttendanceExportAdminTests(TestCase):
    def setUp(self):
        self.superuser = User.objects.create_superuser(
            username='attendance-export-admin',
            email='export-admin@example.com',
            password='test-password',
        )
        self.policy = WorkPolicy.objects.create(
            name='Export Office Policy',
            work_start=time(9, 0),
            work_end=time(17, 0),
            workdays=[0, 1, 2, 3, 4],
        )
        self.first_user = User.objects.create_user(
            username='ayu-export',
            password='test-password',
            is_staff=True,
        )
        self.first_profile = Profile.objects.create(
            user=self.first_user,
            full_name='Ayu Export',
            role=RoleType.ARCHITECT,
            gender=GenderType.FEMALE,
            status=StatusType.PERMANENT,
            birthday=date(1992, 2, 2),
            join_date=date(2024, 2, 2),
            phone_number='0812000001',
            work_policy=self.policy,
        )
        self.second_user = User.objects.create_user(
            username='bima-export',
            password='test-password',
            is_staff=True,
        )
        self.second_profile = Profile.objects.create(
            user=self.second_user,
            full_name='Bima Export',
            role=RoleType.WORKER,
            gender=GenderType.MALE,
            status=StatusType.PERMANENT,
            birthday=date(1991, 1, 1),
            join_date=date(2024, 1, 1),
            phone_number='0812000002',
            work_policy=self.policy,
        )
        self.first_attendance = Attendance.objects.create(
            user=self.first_profile,
            date=date(2026, 7, 28),
            check_in=timezone.make_aware(
                datetime(2026, 7, 28, 8, 55)
            ),
            check_out=timezone.make_aware(
                datetime(2026, 7, 28, 17, 15)
            ),
            status=AttendanceStatus.ONTIME,
            work_policy=self.policy,
            work_mode=AttendanceWorkMode.OFFICE,
            check_in_location_label='Office · Denpasar',
        )
        self.second_attendance = Attendance.objects.create(
            user=self.second_profile,
            date=date(2026, 7, 29),
            check_in=timezone.make_aware(
                datetime(2026, 7, 29, 9, 20)
            ),
            status=AttendanceStatus.LATE,
            work_policy=self.policy,
            work_mode=AttendanceWorkMode.OFFICE,
        )
        self.report_url = reverse(
            'admin:team_attendance_report'
        )
        self.excel_url = reverse(
            'admin:team_attendance_export_excel'
        )
        self.pdf_url = reverse(
            'admin:team_attendance_export_pdf'
        )

    def test_excel_report_filters_requested_date_range_and_has_summary(self):
        self.client.force_login(self.superuser)

        response = self.client.post(
            self.report_url,
            {
                'date_from': '2026-07-28',
                'date_to': '2026-07-28',
                'output_format': 'xlsx',
            },
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            (
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'
            ),
        )
        self.assertIn(
            'attachment; filename="attendance_report_20260728-20260728_',
            response['Content-Disposition'],
        )
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            ['Attendance', 'Ringkasan', 'Per Staff'],
        )
        self.assertEqual(len(workbook['Attendance']._images), 1)
        self.assertEqual(len(workbook['Ringkasan']._images), 1)
        self.assertEqual(len(workbook['Per Staff']._images), 1)
        worksheet = workbook['Attendance']
        self.assertEqual(worksheet['B5'].value, 'Ayu Export')
        self.assertIsNone(worksheet['B6'].value)
        self.assertEqual(
            workbook['Ringkasan']['B3'].value,
            '28/07/2026 – 28/07/2026',
        )
        self.assertEqual(workbook['Ringkasan']['B4'].value, 1)
        staff_sheet = workbook['Per Staff']
        self.assertEqual(staff_sheet['B5'].value, 'Ayu Export')
        self.assertEqual(staff_sheet['E5'].value, 1)
        self.assertEqual(staff_sheet['F5'].value, 1)
        self.assertEqual(staff_sheet['G5'].value, 1)
        self.assertEqual(staff_sheet['H5'].value, 0)
        self.assertIsNone(staff_sheet['B6'].value)

    def test_excel_report_summarizes_each_staff_attendance_metrics(self):
        self.client.force_login(self.superuser)

        response = self.client.post(
            self.report_url,
            {
                'date_from': '2026-07-28',
                'date_to': '2026-07-29',
                'output_format': 'xlsx',
            },
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        staff_sheet = workbook['Per Staff']
        self.assertEqual(staff_sheet['B5'].value, 'Ayu Export')
        self.assertEqual(staff_sheet['G5'].value, 1)
        self.assertEqual(staff_sheet['H5'].value, 0)
        self.assertEqual(staff_sheet['B6'].value, 'Bima Export')
        self.assertEqual(staff_sheet['E6'].value, 1)
        self.assertEqual(staff_sheet['F6'].value, 1)
        self.assertEqual(staff_sheet['G6'].value, 0)
        self.assertEqual(staff_sheet['H6'].value, 1)
        self.assertEqual(staff_sheet['I6'].value, 0)
        self.assertEqual(staff_sheet['K6'].value, 0)

    def test_pdf_report_filters_requested_date_range(self):
        self.client.force_login(self.superuser)

        response = self.client.post(
            self.report_url,
            {
                'date_from': '2026-07-29',
                'date_to': '2026-07-29',
                'output_format': 'pdf',
            },
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF-1.4'))
        self.assertTrue(response.content.rstrip().endswith(b'%%EOF'))
        self.assertIn(b'/Subtype /Image', response.content)
        self.assertIn(b'/Logo Do', response.content)
        self.assertIn(b'Ringkasan per Staff', response.content)
        self.assertIn(b'Detail Attendance', response.content)
        self.assertIn(b'Bima Export', response.content)
        self.assertNotIn(b'Ayu Export', response.content)

    def test_staff_report_contains_only_own_visible_attendance(self):
        view_permission = Permission.objects.get(
            content_type__app_label='team',
            codename='view_attendance',
        )
        self.first_user.user_permissions.add(view_permission)
        self.client.force_login(self.first_user)

        response = self.client.post(
            self.report_url,
            {
                'date_from': '2026-07-28',
                'date_to': '2026-07-29',
                'output_format': 'xlsx',
            },
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook['Attendance']
        exported_names = [
            cell.value
            for cell in worksheet['B'][4:]
            if cell.value
        ]
        self.assertEqual(exported_names, ['Ayu Export'])

    def test_report_page_rejects_reversed_date_range(self):
        self.client.force_login(self.superuser)

        response = self.client.post(
            self.report_url,
            {
                'date_from': '2026-07-29',
                'date_to': '2026-07-28',
                'output_format': 'xlsx',
            },
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            'Tanggal akhir tidak boleh sebelum tanggal mulai.',
        )

    def test_report_page_defaults_to_current_month(self):
        self.client.force_login(self.superuser)

        response = self.client.get(self.report_url, secure=True)

        today = timezone.localdate()
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Laporan Attendance')
        self.assertContains(
            response,
            f'value="{today.replace(day=1):%Y-%m-%d}"',
            html=False,
        )
        self.assertContains(
            response,
            f'value="{today:%Y-%m-%d}"',
            html=False,
        )

    def test_changelist_shows_single_attendance_report_button(self):
        self.client.force_login(self.superuser)

        response = self.client.get(
            reverse('admin:team_attendance_changelist'),
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Laporan Attendance')
        self.assertContains(response, 'TEAM ATTENDANCE')
        self.assertContains(response, 'Heatmap')
        self.assertContains(response, 'Ayu Export')
        self.assertContains(response, self.report_url)
        self.assertNotContains(response, 'Export Excel')
        self.assertNotContains(response, 'Export PDF')

    def test_legacy_export_urls_redirect_to_report_form(self):
        self.client.force_login(self.superuser)

        excel_response = self.client.get(self.excel_url, secure=True)
        pdf_response = self.client.get(self.pdf_url, secure=True)

        self.assertRedirects(
            excel_response,
            f'{self.report_url}?format=xlsx',
            fetch_redirect_response=False,
        )
        self.assertRedirects(
            pdf_response,
            f'{self.report_url}?format=pdf',
            fetch_redirect_response=False,
        )


class SimplifiedAttendanceWorkflowTests(TestCase):
    def setUp(self):
        self.media_directory = TemporaryDirectory()
        media_override = override_settings(
            MEDIA_ROOT=self.media_directory.name,
        )
        media_override.enable()
        self.addCleanup(media_override.disable)
        self.addCleanup(self.media_directory.cleanup)

        self.location = Location.objects.create(
            name='Kantor Denpasar',
            address=Point(115.2616, -8.6538, srid=4326),
        )
        self.policy = WorkPolicy.objects.create(
            name='Office Attendance',
            office_location=self.location,
            work_start=time(9, 0),
            work_end=time(17, 0),
            late_grace_minutes=10,
            overtime_after_minutes=30,
            geofence_radius_meters=100,
            workdays=[0, 1, 2, 3, 4],
        )
        self.user = User.objects.create_user(
            username='simple-attendance',
            password='test-password',
            is_staff=True,
        )
        self.profile = Profile.objects.create(
            user=self.user,
            location=self.location,
            full_name='Simple Attendance',
            role=RoleType.WORKER,
            gender=GenderType.MALE,
            status=StatusType.PERMANENT,
            birthday=date(1990, 1, 1),
            join_date=date(2024, 1, 1),
            phone_number='0812222222',
            work_policy=self.policy,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.admin_instance = admin.site._registry[Attendance]
        self.request_factory = RequestFactory()

    def image_upload(self, name):
        image = Image.new('RGB', (2, 2), color='white')
        content = BytesIO()
        image.save(content, format='JPEG')
        content.seek(0)
        return SimpleUploadedFile(
            name,
            content.read(),
            content_type='image/jpeg',
        )

    def location_payload(self, point=None):
        point = point or self.location.address
        return json.dumps({
            'type': 'Point',
            'coordinates': [point.x, point.y],
        })

    def test_api_checkin_and_checkout_need_only_visible_fields(self):
        check_in_response = self.client.post(
            '/api/team/checkin/',
            {
                'work_policy': self.policy.pk,
                'work_mode': AttendanceWorkMode.OFFICE,
                'photo_check_in': self.image_upload('check-in.jpg'),
                'check_in_location': self.location_payload(),
                'check_in_accuracy_meters': '8.5',
            },
            format='multipart',
            secure=True,
        )

        self.assertEqual(check_in_response.status_code, 201)
        attendance = Attendance.objects.get(
            user=self.profile,
            date=timezone.localdate(),
        )
        self.assertEqual(attendance.work_policy, self.policy)
        self.assertEqual(
            attendance.check_in_location_label,
            f'Office · {self.location.name}',
        )
        self.assertAlmostEqual(
            attendance.check_in_location.x,
            self.location.address.x,
        )
        self.assertIsNotNone(attendance.check_in)
        self.assertEqual(attendance.check_in_accuracy_meters, 8.5)
        self.assertAlmostEqual(
            attendance.check_in_distance_meters,
            0,
            places=2,
        )

        check_out_response = self.client.post(
            '/api/team/checkout/',
            {
                'photo_check_out': self.image_upload('check-out.jpg'),
                'check_out_location': self.location_payload(),
                'check_out_accuracy_meters': '9.25',
            },
            format='multipart',
            secure=True,
        )

        self.assertEqual(check_out_response.status_code, 201)
        attendance.refresh_from_db()
        self.assertIsNotNone(attendance.check_out)
        self.assertEqual(
            attendance.check_out_location_label,
            f'Office · {self.location.name}',
        )
        self.assertEqual(attendance.check_out_accuracy_meters, 9.25)

    def test_api_rejects_missing_or_out_of_range_live_location(self):
        missing_location = self.client.post(
            '/api/team/checkin/',
            {
                'work_policy': self.policy.pk,
                'photo_check_in': self.image_upload('missing-gps.jpg'),
            },
            format='multipart',
            secure=True,
        )

        self.assertEqual(missing_location.status_code, 400)
        self.assertIn('Lokasi GPS terkini wajib', missing_location.data['error'])

        far_location = Point(115.30, -8.70, srid=4326)
        outside_geofence = self.client.post(
            '/api/team/checkin/',
            {
                'work_policy': self.policy.pk,
                'photo_check_in': self.image_upload('far-away.jpg'),
                'check_in_location': self.location_payload(far_location),
            },
            format='multipart',
            secure=True,
        )

        self.assertEqual(outside_geofence.status_code, 400)
        self.assertIn('berjarak', outside_geofence.data['error'])
        self.assertIn('Batas untuk Office adalah 100 meter', outside_geofence.data['error'])
        self.assertFalse(
            Attendance.objects.filter(
                user=self.profile,
                date=timezone.localdate(),
            ).exists()
        )

    def test_wfh_requires_permission_and_uses_approved_home_location(self):
        home = Location.objects.create(
            name='Rumah Staff',
            address=Point(115.20, -8.60, srid=4326),
        )
        self.profile.location = home
        self.profile.save(update_fields=('location',))

        rejected = self.client.post(
            '/api/team/checkin/',
            {
                'work_policy': self.policy.pk,
                'work_mode': AttendanceWorkMode.WFH,
                'photo_check_in': self.image_upload('wfh-rejected.jpg'),
                'check_in_location': self.location_payload(home.address),
            },
            format='multipart',
            secure=True,
        )

        self.assertEqual(rejected.status_code, 400)
        self.assertIn('tidak diizinkan', rejected.data['error'])

        self.policy.allow_wfh = True
        self.policy.wfh_geofence_radius_meters = 150
        self.policy.save(
            update_fields=(
                'allow_wfh',
                'wfh_geofence_radius_meters',
            )
        )
        accepted = self.client.post(
            '/api/team/checkin/',
            {
                'work_policy': self.policy.pk,
                'work_mode': AttendanceWorkMode.WFH,
                'photo_check_in': self.image_upload('wfh-accepted.jpg'),
                'check_in_location': self.location_payload(home.address),
                'check_in_accuracy_meters': '12',
            },
            format='multipart',
            secure=True,
        )

        self.assertEqual(accepted.status_code, 201)
        attendance = Attendance.objects.get(
            user=self.profile,
            date=timezone.localdate(),
        )
        self.assertEqual(attendance.work_mode, AttendanceWorkMode.WFH)
        self.assertEqual(
            attendance.check_in_location_label,
            f'WFH · {home.name}',
        )
        self.assertAlmostEqual(
            attendance.check_in_location.x,
            home.address.x,
        )

    def test_generic_attendance_mutation_is_superuser_only(self):
        response = self.client.post(
            '/api/team/attendance/',
            {
                'user_id': self.profile.pk,
                'date': timezone.localdate(),
                'work_policy': self.policy.pk,
            },
            format='json',
            secure=True,
        )

        self.assertEqual(response.status_code, 403)

    def test_non_superuser_admin_has_minimal_fields(self):
        request = self.request_factory.get('/admin/team/attendance/add/')
        request.user = self.user

        add_fieldsets = self.admin_instance.get_fieldsets(request)
        self.assertEqual(
            add_fieldsets[0][1]['fields'],
            (
                'work_policy',
                'work_mode',
                'photo_check_in',
                'check_in_location',
                'check_in_gps_accuracy',
            ),
        )

        form_class = self.admin_instance.get_form(request)
        form = form_class(
            data={
                'work_policy': self.policy.pk,
                'work_mode': AttendanceWorkMode.OFFICE,
                'check_in_location': (
                    f'SRID=4326;POINT ({self.location.address.x} '
                    f'{self.location.address.y})'
                ),
                'check_in_gps_accuracy': '7',
            },
            files={
                'photo_check_in': self.image_upload('admin-in.jpg'),
            },
        )
        self.assertTrue(form.is_valid(), form.errors)
        attendance = form.save(commit=False)
        self.admin_instance.save_model(
            request,
            attendance,
            form,
            change=False,
        )

        self.assertEqual(attendance.user, self.profile)
        self.assertEqual(attendance.date, timezone.localdate())
        self.assertIsNotNone(attendance.check_in)
        self.assertEqual(
            attendance.check_in_location_label,
            f'Office · {self.location.name}',
        )
        self.assertEqual(attendance.check_in_accuracy_meters, 7)

        checkout_request = self.request_factory.post(
            f'/admin/team/attendance/{attendance.pk}/change/'
        )
        checkout_request.user = self.user
        checkout_fieldsets = self.admin_instance.get_fieldsets(
            checkout_request,
            attendance,
        )
        self.assertEqual(
            checkout_fieldsets[1][1]['fields'],
            (
                'photo_check_out',
                'check_out_location',
                'check_out_gps_accuracy',
            ),
        )

    def test_non_superuser_admin_add_page_renders_only_checkin_inputs(self):
        project_location = Location.objects.create(
            name='Proyek Lapangan',
            address=Point(115.2716, -8.6638, srid=4326),
        )
        field_policy = WorkPolicy.objects.create(
            name='Dinas Proyek Lapangan',
            office_location=project_location,
            work_start=time(8, 0),
            work_end=time(17, 0),
            geofence_radius_meters=200,
        )
        WorkPolicy.objects.create(
            name='Policy Tidak Aktif',
            office_location=project_location,
            work_start=time(8, 0),
            work_end=time(17, 0),
            is_active=False,
        )
        self.user.user_permissions.add(
            *Permission.objects.filter(
                content_type__app_label='team',
                codename__in=[
                    'add_attendance',
                    'view_attendance',
                    'change_attendance',
                ],
            )
        )
        admin_client = APIClient()
        admin_client.force_login(self.user)

        response = admin_client.get(
            '/admin/team/attendance/add/',
            secure=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id_work_policy')
        self.assertContains(response, self.policy.name)
        self.assertContains(response, field_policy.name)
        self.assertNotContains(response, 'Policy Tidak Aktif')
        self.assertContains(
            response,
            'Pilih lokasi kerja Anda saat ini',
        )
        self.assertContains(response, 'id_work_mode')
        self.assertContains(response, 'id_photo_check_in')
        self.assertContains(response, 'id_check_in_location')
        self.assertContains(response, 'id_check_in_gps_accuracy')
        self.assertContains(
            response,
            'data-mmg-camera-required="true"',
        )
        self.assertContains(response, 'capture="user"')
        self.assertContains(
            response,
            static('admin/css/attendance_camera.css'),
        )
        self.assertContains(
            response,
            static('admin/js/attendance_camera.js'),
        )
        self.assertContains(
            response,
            'ambil foto langsung dari kamera',
        )
        self.assertNotContains(response, 'id_date')
        self.assertNotContains(response, 'id_photo_check_out')

    def test_superuser_admin_can_correct_operational_fields(self):
        superuser = User.objects.create_superuser(
            username='attendance-superuser',
            email='superuser@example.com',
            password='test-password',
        )
        request = self.request_factory.get(
            '/admin/team/attendance/add/'
        )
        request.user = superuser

        fieldsets = self.admin_instance.get_fieldsets(request)
        visible_fields = {
            field
            for _, options in fieldsets
            for row in options['fields']
            for field in (row if isinstance(row, tuple) else (row,))
        }
        readonly_fields = self.admin_instance.get_readonly_fields(
            request,
        )
        form = self.admin_instance.get_form(request)()
        photo_widget = form.fields['photo_check_in'].widget

        self.assertIn('date', visible_fields)
        self.assertIn('check_in', visible_fields)
        self.assertIn('check_out', visible_fields)
        self.assertIn('check_in_location', visible_fields)
        self.assertIn('work_mode', visible_fields)
        self.assertIn('status_override', visible_fields)
        self.assertNotIn('date', readonly_fields)
        self.assertNotIn('check_in', readonly_fields)
        self.assertNotIn('status_override', readonly_fields)
        self.assertIn(
            'mmg-styled-file-input',
            photo_widget.attrs['class'],
        )
        self.assertNotIn(
            'data-mmg-camera-required',
            photo_widget.attrs,
        )


class ImageCompressionTests(TestCase):
    def setUp(self):
        self.media_directory = TemporaryDirectory()
        media_override = override_settings(
            MEDIA_ROOT=self.media_directory.name,
            IMAGE_UPLOAD_MAX_WIDTH=800,
            IMAGE_UPLOAD_MAX_HEIGHT=800,
            IMAGE_UPLOAD_QUALITY=75,
        )
        media_override.enable()
        self.addCleanup(media_override.disable)
        self.addCleanup(self.media_directory.cleanup)

    def test_audit_model_compresses_new_image_before_storage(self):
        image = Image.new('RGBA', (2400, 1200), color=(20, 30, 40, 180))
        content = BytesIO()
        image.save(content, format='PNG')
        upload = SimpleUploadedFile(
            'large-profile.png',
            content.getvalue(),
            content_type='image/png',
        )
        user = User.objects.create_user(
            username='compressed-profile',
            password='test-password',
        )

        profile = Profile.objects.create(
            user=user,
            full_name='Compressed Profile',
            role=RoleType.WORKER,
            gender=GenderType.MALE,
            status=StatusType.PERMANENT,
            birthday=date(1990, 1, 1),
            join_date=date(2024, 1, 1),
            phone_number='0812999999',
            profile_picture=upload,
        )

        with profile.profile_picture.open('rb') as stored_file:
            with Image.open(stored_file) as stored_image:
                self.assertEqual(stored_image.format, 'JPEG')
                self.assertLessEqual(stored_image.width, 800)
                self.assertLessEqual(stored_image.height, 800)
                self.assertEqual(stored_image.mode, 'RGB')
