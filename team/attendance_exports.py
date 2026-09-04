"""Excel and PDF exports for attendance records.

The PDF writer intentionally uses PDF's built-in Type 1 fonts so attendance
exports do not require a native/system PDF dependency.
"""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Iterable, Sequence
import zlib

from django.conf import settings
from django.contrib.staticfiles import finders
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

from .models import Attendance, AttendanceStatus, AttendanceWorkMode


GOLD = "D4AF37"
BLACK = "171717"
WHITE = "FFFFFF"
LIGHT_GOLD = "F7F0D8"
LIGHT_GRAY = "F3F4F6"
MID_GRAY = "D1D5DB"
REPORT_LOGO_STATIC_PATH = "favicon/MMG_PNG_WHITE.png"


def _report_logo_assets() -> dict | None:
    """Buat varian logo emas transparan (Excel) dan RGB (PDF)."""
    located_path = finders.find(REPORT_LOGO_STATIC_PATH)
    candidates = [
        Path(located_path) if located_path else None,
        Path(settings.BASE_DIR)
        / "admin-interface"
        / REPORT_LOGO_STATIC_PATH,
        Path(settings.BASE_DIR) / "static" / REPORT_LOGO_STATIC_PATH,
    ]
    source_path = next(
        (
            candidate
            for candidate in candidates
            if candidate is not None and candidate.is_file()
        ),
        None,
    )
    if source_path is None:
        return None

    try:
        with PILImage.open(source_path) as source:
            rgba = source.convert("RGBA")
            alpha = rgba.getchannel("A")
            bounding_box = alpha.getbbox()
            if bounding_box is None:
                return None
            alpha = alpha.crop(bounding_box)
            logo = PILImage.new(
                "RGBA",
                alpha.size,
                (212, 175, 55, 0),
            )
            logo.putalpha(alpha)

            png_output = BytesIO()
            logo.save(png_output, format="PNG", optimize=True)

            pdf_logo = PILImage.new(
                "RGB",
                logo.size,
                tuple(int(BLACK[index:index + 2], 16) for index in (0, 2, 4)),
            )
            pdf_logo.paste(logo, mask=logo.getchannel("A"))
            return {
                "png": png_output.getvalue(),
                "pdf_rgb": zlib.compress(pdf_logo.tobytes()),
                "width": logo.width,
                "height": logo.height,
            }
    except (OSError, ValueError):
        return None


def _add_excel_logo(
    worksheet,
    logo_assets: dict | None,
    anchor: str = "A1",
):
    if not logo_assets:
        return
    logo = ExcelImage(BytesIO(logo_assets["png"]))
    max_width = 39
    max_height = 30
    scale = min(
        max_width / logo.width,
        max_height / logo.height,
    )
    logo.width *= scale
    logo.height *= scale
    worksheet.add_image(logo, anchor)


def _safe_excel_text(value) -> str:
    """Prevent spreadsheet formula injection in user-controlled text."""
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _local_datetime(value):
    if value is None:
        return None
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.replace(tzinfo=None)


def _display_minutes(value) -> str:
    minutes = max(int(value or 0), 0)
    hours, remainder = divmod(minutes, 60)
    return f"{hours}j {remainder:02d}m"


def _display_datetime(value, empty="-") -> str:
    local_value = _local_datetime(value)
    if local_value is None:
        return empty
    return local_value.strftime("%d/%m/%Y %H:%M")


def _display_time(value, empty="-") -> str:
    local_value = _local_datetime(value)
    if local_value is None:
        return empty
    return local_value.strftime("%H:%M")


def _display_point(point) -> str:
    if not point:
        return "-"
    return f"{point.y:.6f}, {point.x:.6f}"


def _display_number(value):
    if value is None:
        return None
    return round(float(value), 2)


def _photo_status(field) -> str:
    return "Ada" if field and getattr(field, "name", "") else "Tidak ada"


def _records_summary(
    records: Sequence[Attendance],
    date_from: date | None = None,
    date_to: date | None = None,
) -> dict:
    dates = [record.date for record in records if record.date]
    users = {record.user_id for record in records}
    statuses = Counter(record.get_status_display() for record in records)
    work_modes = Counter(record.get_work_mode_display() for record in records)
    return {
        "count": len(records),
        "staff_count": len(users),
        "date_from": date_from or (min(dates) if dates else None),
        "date_to": date_to or (max(dates) if dates else None),
        "worked_minutes": sum(record.worked_minutes or 0 for record in records),
        "overtime_minutes": sum(
            record.overtime_minutes or 0 for record in records
        ),
        "statuses": statuses,
        "work_modes": work_modes,
    }


def _staff_summaries(records: Sequence[Attendance]) -> list[dict]:
    """Agregasi laporan per staff tanpa menghilangkan status gabungan."""
    summaries = {}
    present_statuses = {
        AttendanceStatus.ONTIME,
        AttendanceStatus.LATE,
        AttendanceStatus.EARLY_LEAVE,
        AttendanceStatus.LATE_EARLY_LEAVE,
        AttendanceStatus.OVERTIME,
    }
    ontime_statuses = {
        AttendanceStatus.ONTIME,
        AttendanceStatus.OVERTIME,
    }
    late_statuses = {
        AttendanceStatus.LATE,
        AttendanceStatus.LATE_EARLY_LEAVE,
    }
    early_leave_statuses = {
        AttendanceStatus.EARLY_LEAVE,
        AttendanceStatus.LATE_EARLY_LEAVE,
    }

    for record in records:
        profile = record.user
        auth_user = profile.user
        summary = summaries.setdefault(
            record.user_id,
            {
                "name": profile.full_name,
                "username": auth_user.username,
                "role": profile.get_role_display(),
                "days": 0,
                "present": 0,
                "ontime": 0,
                "late": 0,
                "early_leave": 0,
                "overtime": 0,
                "absent": 0,
                "leave": 0,
                "holiday": 0,
                "wfh": 0,
                "worked_minutes": 0,
                "overtime_minutes": 0,
            },
        )
        status = record.status
        summary["days"] += 1
        summary["present"] += int(
            bool(record.check_in) or status in present_statuses
        )
        summary["ontime"] += int(status in ontime_statuses)
        summary["late"] += int(status in late_statuses)
        summary["early_leave"] += int(status in early_leave_statuses)
        summary["overtime"] += int(
            status == AttendanceStatus.OVERTIME
            or bool(record.overtime_minutes)
        )
        summary["absent"] += int(status == AttendanceStatus.ABSENT)
        summary["leave"] += int(status == AttendanceStatus.LEAVE)
        summary["holiday"] += int(status == AttendanceStatus.HOLYDAY)
        summary["wfh"] += int(
            record.work_mode == AttendanceWorkMode.WFH
        )
        summary["worked_minutes"] += record.worked_minutes or 0
        summary["overtime_minutes"] += record.overtime_minutes or 0

    return sorted(
        summaries.values(),
        key=lambda item: (
            str(item["name"]).casefold(),
            str(item["username"]).casefold(),
        ),
    )


def attendance_export_filename(
    records: Sequence[Attendance],
    extension: str,
    generated_at=None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> str:
    generated_at = _local_datetime(generated_at or timezone.now())
    dates = [record.date for record in records if record.date]
    report_date_from = date_from or (min(dates) if dates else None)
    report_date_to = date_to or (max(dates) if dates else None)
    if report_date_from and report_date_to:
        period = f"{report_date_from:%Y%m%d}-{report_date_to:%Y%m%d}"
    else:
        period = "tanpa-data"
    return (
        f"attendance_report_{period}_"
        f"{generated_at:%Y%m%d_%H%M}.{extension}"
    )


def build_attendance_excel(
    attendances: Iterable[Attendance],
    generated_at=None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> bytes:
    records = list(attendances)
    generated_at = generated_at or timezone.now()
    summary = _records_summary(
        records,
        date_from=date_from,
        date_to=date_to,
    )
    staff_summaries = _staff_summaries(records)
    logo_assets = _report_logo_assets()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Attendance"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A5"

    headers = [
        "No.",
        "Nama staff",
        "Username",
        "Role",
        "Tanggal",
        "Work policy",
        "Mode kerja",
        "Status",
        "Check-in",
        "Check-out",
        "Jam kerja",
        "Lembur",
        "Lokasi check-in",
        "Jarak check-in (m)",
        "Akurasi check-in (m)",
        "Koordinat check-in",
        "Lokasi check-out",
        "Jarak check-out (m)",
        "Akurasi check-out (m)",
        "Koordinat check-out",
        "Foto check-in",
        "Foto check-out",
        "Status dikoreksi",
        "Alasan koreksi",
    ]

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(headers),
    )
    title_cell = worksheet.cell(
        1,
        1,
        "MMG CONSTRUCTION · LAPORAN ATTENDANCE",
    )
    title_cell.font = Font(color=GOLD, bold=True, size=16)
    title_cell.fill = PatternFill("solid", fgColor=BLACK)
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    worksheet.row_dimensions[1].height = 30
    _add_excel_logo(worksheet, logo_assets)

    period = "-"
    if summary["date_from"]:
        period = (
            f"{summary['date_from']:%d/%m/%Y} – "
            f"{summary['date_to']:%d/%m/%Y}"
        )
    generated_local = _local_datetime(generated_at)
    worksheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=len(headers),
    )
    worksheet.cell(
        2,
        1,
        (
            f"Periode: {period}  |  {summary['count']} data  |  "
            f"{summary['staff_count']} staff  |  "
            f"Dibuat: {generated_local:%d/%m/%Y %H:%M}"
        ),
    )
    worksheet.cell(2, 1).font = Font(color="4B5563", italic=True, size=10)
    worksheet.cell(2, 1).alignment = Alignment(vertical="center")
    worksheet.row_dimensions[2].height = 22

    header_row = 4
    thin_border = Border(
        left=Side(style="thin", color=MID_GRAY),
        right=Side(style="thin", color=MID_GRAY),
        top=Side(style="thin", color=MID_GRAY),
        bottom=Side(style="thin", color=MID_GRAY),
    )
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(header_row, column, header)
        cell.font = Font(color=BLACK, bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor=GOLD)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border
    worksheet.row_dimensions[header_row].height = 34

    for row_number, record in enumerate(records, start=1):
        excel_row = header_row + row_number
        profile = record.user
        auth_user = profile.user
        values = [
            row_number,
            _safe_excel_text(profile.full_name),
            _safe_excel_text(auth_user.username),
            _safe_excel_text(profile.get_role_display()),
            record.date,
            _safe_excel_text(record.work_policy.name)
            if record.work_policy
            else "-",
            record.get_work_mode_display(),
            record.get_status_display(),
            _local_datetime(record.check_in),
            _local_datetime(record.check_out),
            round((record.worked_minutes or 0) / 60, 2),
            round((record.overtime_minutes or 0) / 60, 2),
            _safe_excel_text(record.check_in_location_label or "-"),
            _display_number(record.check_in_distance_meters),
            _display_number(record.check_in_accuracy_meters),
            _display_point(record.check_in_location),
            _safe_excel_text(record.check_out_location_label or "-"),
            _display_number(record.check_out_distance_meters),
            _display_number(record.check_out_accuracy_meters),
            _display_point(record.check_out_location),
            _photo_status(record.photo_check_in),
            _photo_status(record.photo_check_out),
            record.get_status_override_display()
            if record.status_override
            else "-",
            _safe_excel_text(record.status_override_reason or "-"),
        ]
        fill_color = WHITE if row_number % 2 else LIGHT_GRAY
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(excel_row, column, value)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column in {2, 6, 13, 17, 24},
            )
            cell.font = Font(color=BLACK, size=9)

        worksheet.cell(excel_row, 5).number_format = "dd/mm/yyyy"
        worksheet.cell(excel_row, 9).number_format = "dd/mm/yyyy hh:mm"
        worksheet.cell(excel_row, 10).number_format = "dd/mm/yyyy hh:mm"
        for column in (11, 12):
            worksheet.cell(excel_row, column).number_format = '0.00 "jam"'
        for column in (14, 15, 18, 19):
            worksheet.cell(excel_row, column).number_format = '0.00'
        worksheet.row_dimensions[excel_row].height = 30

    last_row = max(header_row, header_row + len(records))
    worksheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
    )
    widths = [
        6, 25, 18, 18, 13, 23, 17, 20, 19, 19, 13, 13,
        35, 18, 19, 24, 35, 19, 20, 24, 16, 16, 20, 40,
    ]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = f"1:{header_row}"
    worksheet.print_area = (
        f"A1:{get_column_letter(len(headers))}{last_row}"
    )

    summary_sheet = workbook.create_sheet("Ringkasan")
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.column_dimensions["A"].width = 30
    summary_sheet.column_dimensions["B"].width = 30
    summary_sheet.merge_cells("A1:B1")
    summary_sheet["A1"] = "RINGKASAN ATTENDANCE"
    summary_sheet["A1"].fill = PatternFill("solid", fgColor=BLACK)
    summary_sheet["A1"].font = Font(color=GOLD, bold=True, size=15)
    summary_sheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    summary_sheet.row_dimensions[1].height = 30
    _add_excel_logo(summary_sheet, logo_assets)

    summary_rows = [
        ("Periode", period),
        ("Total data", summary["count"]),
        ("Jumlah staff", summary["staff_count"]),
        ("Total jam kerja", round(summary["worked_minutes"] / 60, 2)),
        ("Total jam lembur", round(summary["overtime_minutes"] / 60, 2)),
        ("Dibuat", generated_local),
    ]
    current_row = 3
    for label, value in summary_rows:
        summary_sheet.cell(current_row, 1, label)
        summary_sheet.cell(current_row, 2, value)
        summary_sheet.cell(current_row, 1).font = Font(bold=True)
        summary_sheet.cell(current_row, 1).fill = PatternFill(
            "solid", fgColor=LIGHT_GOLD
        )
        for column in (1, 2):
            summary_sheet.cell(current_row, column).border = thin_border
        current_row += 1
    summary_sheet.cell(8, 2).number_format = "dd/mm/yyyy hh:mm"

    for section, counts in (
        ("STATUS", summary["statuses"]),
        ("MODE KERJA", summary["work_modes"]),
    ):
        current_row += 1
        summary_sheet.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=2,
        )
        section_cell = summary_sheet.cell(current_row, 1, section)
        section_cell.fill = PatternFill("solid", fgColor=GOLD)
        section_cell.font = Font(color=BLACK, bold=True)
        current_row += 1
        if not counts:
            counts = {"Tidak ada data": 0}
        for label, count in sorted(counts.items()):
            summary_sheet.cell(current_row, 1, label)
            summary_sheet.cell(current_row, 2, count)
            for column in (1, 2):
                summary_sheet.cell(current_row, column).border = thin_border
            current_row += 1

    staff_sheet = workbook.create_sheet("Per Staff")
    staff_sheet.sheet_view.showGridLines = False
    staff_sheet.freeze_panes = "A5"
    staff_headers = [
        "No.",
        "Nama staff",
        "Username",
        "Role",
        "Hari tercatat",
        "Hadir",
        "Tepat waktu",
        "Terlambat",
        "Pulang cepat",
        "Lembur",
        "Absent",
        "Cuti",
        "Libur",
        "WFH",
        "Total jam kerja",
        "Total jam lembur",
    ]
    staff_sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(staff_headers),
    )
    staff_title = staff_sheet.cell(
        1,
        1,
        "MMG CONSTRUCTION · RINGKASAN ATTENDANCE PER STAFF",
    )
    staff_title.font = Font(color=GOLD, bold=True, size=15)
    staff_title.fill = PatternFill("solid", fgColor=BLACK)
    staff_title.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    staff_sheet.row_dimensions[1].height = 30
    _add_excel_logo(staff_sheet, logo_assets)

    staff_sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=len(staff_headers),
    )
    staff_sheet.cell(
        2,
        1,
        (
            f"Periode: {period}  |  "
            f"{summary['staff_count']} staff  |  "
            f"Dibuat: {generated_local:%d/%m/%Y %H:%M}"
        ),
    )
    staff_sheet.cell(2, 1).font = Font(
        color="4B5563",
        italic=True,
        size=10,
    )

    staff_header_row = 4
    for column, header in enumerate(staff_headers, start=1):
        cell = staff_sheet.cell(staff_header_row, column, header)
        cell.font = Font(color=BLACK, bold=True, size=9)
        cell.fill = PatternFill("solid", fgColor=GOLD)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border
    staff_sheet.row_dimensions[staff_header_row].height = 36

    for row_number, staff in enumerate(staff_summaries, start=1):
        excel_row = staff_header_row + row_number
        values = [
            row_number,
            _safe_excel_text(staff["name"]),
            _safe_excel_text(staff["username"]),
            _safe_excel_text(staff["role"]),
            staff["days"],
            staff["present"],
            staff["ontime"],
            staff["late"],
            staff["early_leave"],
            staff["overtime"],
            staff["absent"],
            staff["leave"],
            staff["holiday"],
            staff["wfh"],
            round(staff["worked_minutes"] / 60, 2),
            round(staff["overtime_minutes"] / 60, 2),
        ]
        fill_color = WHITE if row_number % 2 else LIGHT_GRAY
        for column, value in enumerate(values, start=1):
            cell = staff_sheet.cell(excel_row, column, value)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.border = thin_border
            cell.font = Font(color=BLACK, size=9)
            cell.alignment = Alignment(
                horizontal="left" if column in {2, 3, 4} else "center",
                vertical="center",
                wrap_text=column in {2, 4},
            )
        for column in (15, 16):
            staff_sheet.cell(excel_row, column).number_format = (
                '0.00 "jam"'
            )
        staff_sheet.row_dimensions[excel_row].height = 28

    if not staff_summaries:
        staff_sheet.merge_cells(
            start_row=5,
            start_column=1,
            end_row=5,
            end_column=len(staff_headers),
        )
        empty_cell = staff_sheet.cell(
            5,
            1,
            "Tidak ada data attendance pada periode yang dipilih.",
        )
        empty_cell.alignment = Alignment(horizontal="center")
        empty_cell.font = Font(color="6B7280", italic=True)

    staff_last_row = max(
        staff_header_row,
        staff_header_row + len(staff_summaries),
    )
    staff_sheet.auto_filter.ref = (
        f"A{staff_header_row}:"
        f"{get_column_letter(len(staff_headers))}{staff_last_row}"
    )
    staff_widths = [
        6, 25, 18, 18, 14, 10, 14, 12,
        14, 10, 10, 10, 10, 10, 18, 19,
    ]
    for index, width in enumerate(staff_widths, start=1):
        staff_sheet.column_dimensions[
            get_column_letter(index)
        ].width = width
    staff_sheet.page_setup.orientation = "landscape"
    staff_sheet.page_setup.fitToWidth = 1
    staff_sheet.page_setup.fitToHeight = 0
    staff_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    staff_sheet.print_title_rows = (
        f"1:{staff_header_row}"
    )
    staff_sheet.print_area = (
        f"A1:{get_column_letter(len(staff_headers))}"
        f"{max(5, staff_last_row)}"
    )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_escape(value) -> str:
    text = str(value).encode("latin-1", errors="replace").decode("latin-1")
    return (
        text.replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
        .replace("\r", " ")
        .replace("\n", " ")
    )


def _pdf_text(
    commands: list[str],
    x: float,
    y: float,
    text,
    size: float = 8,
    bold: bool = False,
    color=(0.09, 0.09, 0.09),
):
    font = "F2" if bold else "F1"
    commands.append(
        "BT "
        f"/{font} {size:.1f} Tf "
        f"{color[0]:.3f} {color[1]:.3f} {color[2]:.3f} rg "
        f"1 0 0 1 {x:.2f} {y:.2f} Tm "
        f"({_pdf_escape(text)}) Tj ET"
    )


def _pdf_rect(
    commands: list[str],
    x: float,
    y: float,
    width: float,
    height: float,
    fill,
    stroke=(0.82, 0.82, 0.82),
):
    commands.append(
        f"{fill[0]:.3f} {fill[1]:.3f} {fill[2]:.3f} rg "
        f"{stroke[0]:.3f} {stroke[1]:.3f} {stroke[2]:.3f} RG "
        f"0.4 w {x:.2f} {y:.2f} {width:.2f} {height:.2f} re B"
    )


def _pdf_image(
    commands: list[str],
    x: float,
    y: float,
    width: float,
    height: float,
    name: str = "Logo",
):
    commands.append(
        "q "
        f"{width:.2f} 0 0 {height:.2f} {x:.2f} {y:.2f} cm "
        f"/{name} Do Q"
    )


def _truncate_pdf(value, width: float, font_size: float = 7.2) -> str:
    text = str(value)
    max_characters = max(int(width / (font_size * 0.52)), 1)
    if len(text) <= max_characters:
        return text
    if max_characters <= 3:
        return text[:max_characters]
    return f"{text[:max_characters - 3]}..."


def _build_pdf_document(
    page_streams: Sequence[bytes],
    logo_assets: dict | None = None,
) -> bytes:
    # Object 1: catalog, 2: pages, 3-4: built-in fonts. Each page then
    # occupies two objects (page dictionary + content stream). The optional
    # logo is stored once and reused by every page.
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # Filled after page object numbers are known.
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
    ]
    logo_object_number = None
    if logo_assets:
        logo_object_number = len(objects) + 1
        logo_stream = logo_assets["pdf_rgb"]
        objects.append(
            (
                "<< /Type /XObject /Subtype /Image "
                f"/Width {logo_assets['width']} "
                f"/Height {logo_assets['height']} "
                "/ColorSpace /DeviceRGB /BitsPerComponent 8 "
                "/Filter /FlateDecode "
                f"/Length {len(logo_stream)} >>\nstream\n"
            ).encode("ascii")
            + logo_stream
            + b"\nendstream"
        )

    page_object_numbers = []
    for stream in page_streams:
        page_number = len(objects) + 1
        content_number = page_number + 1
        page_object_numbers.append(page_number)
        logo_resource = ""
        if logo_object_number is not None:
            logo_resource = (
                f"/XObject << /Logo {logo_object_number} 0 R >> "
            )
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R "
                "/MediaBox [0 0 842 595] "
                "/Resources << "
                "/Font << /F1 3 0 R /F2 4 0 R >> "
                f"{logo_resource}>> "
                f"/Contents {content_number} 0 R >>"
            ).encode("ascii")
        )
        objects.append(
            (
                f"<< /Length {len(stream)} >>\nstream\n".encode("ascii")
                + stream
                + b"\nendstream"
            )
        )

    kids = " ".join(f"{number} 0 R" for number in page_object_numbers)
    objects[1] = (
        f"<< /Type /Pages /Kids [{kids}] /Count {len(page_streams)} >>"
    ).encode("ascii")

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(obj)
        output.extend(b"\nendobj\n")

    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            "trailer\n"
            f"<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(output)


def build_attendance_pdf(
    attendances: Iterable[Attendance],
    generated_at=None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> bytes:
    records = list(attendances)
    generated_at = generated_at or timezone.now()
    generated_local = _local_datetime(generated_at)
    summary = _records_summary(
        records,
        date_from=date_from,
        date_to=date_to,
    )
    staff_summaries = _staff_summaries(records)
    logo_assets = _report_logo_assets()

    detail_columns = [
        ("No", 24),
        ("Tanggal", 58),
        ("Staff", 120),
        ("Role", 62),
        ("Mode", 54),
        ("Masuk", 46),
        ("Pulang", 46),
        ("Kerja", 48),
        ("Lembur", 48),
        ("Status", 92),
        ("Work policy", 132),
    ]
    staff_columns = [
        ("No", 24),
        ("Staff", 130),
        ("Role", 60),
        ("Hari", 34),
        ("Hadir", 36),
        ("Tepat", 42),
        ("Telat", 36),
        ("Pulang cepat", 54),
        ("Lembur", 42),
        ("Absent", 40),
        ("Cuti", 34),
        ("Libur", 34),
        ("WFH", 34),
        ("Jam kerja", 52),
        ("Jam lembur", 52),
    ]
    margin_x = 36
    page_height = 595
    detail_row_height = 18
    detail_rows_per_page = 21
    detail_chunks = [
        records[index:index + detail_rows_per_page]
        for index in range(0, len(records), detail_rows_per_page)
    ] or [[]]
    staff_row_height = 19
    staff_rows_per_page = 18
    staff_chunks = [
        staff_summaries[index:index + staff_rows_per_page]
        for index in range(0, len(staff_summaries), staff_rows_per_page)
    ] or [[]]
    total_pages = len(staff_chunks) + len(detail_chunks)
    page_streams: list[bytes] = []

    period = "-"
    if summary["date_from"]:
        period = (
            f"{summary['date_from']:%d/%m/%Y} - "
            f"{summary['date_to']:%d/%m/%Y}"
        )

    def add_page_header(commands, subtitle):
        _pdf_rect(
            commands,
            0,
            page_height - 72,
            842,
            72,
            fill=(0.09, 0.09, 0.09),
            stroke=(0.09, 0.09, 0.09),
        )
        _pdf_rect(
            commands,
            0,
            page_height - 76,
            842,
            4,
            fill=(0.831, 0.686, 0.216),
            stroke=(0.831, 0.686, 0.216),
        )
        title_x = margin_x
        if logo_assets:
            logo_width = 42
            logo_height = (
                logo_width
                * logo_assets["height"]
                / logo_assets["width"]
            )
            _pdf_image(
                commands,
                margin_x,
                page_height - 63,
                logo_width,
                logo_height,
            )
            title_x = margin_x + 53
        _pdf_text(
            commands,
            title_x,
            page_height - 34,
            "MMG CONSTRUCTION",
            size=17,
            bold=True,
            color=(0.831, 0.686, 0.216),
        )
        _pdf_text(
            commands,
            title_x,
            page_height - 54,
            subtitle,
            size=10,
            color=(1, 1, 1),
        )
        _pdf_text(
            commands,
            610,
            page_height - 34,
            f"Periode {period}",
            size=8,
            bold=True,
            color=(1, 1, 1),
        )
        _pdf_text(
            commands,
            610,
            page_height - 51,
            (
                f"{summary['count']} data · "
                f"{summary['staff_count']} staff"
            ),
            size=8,
            color=(0.88, 0.88, 0.88),
        )

    def add_page_footer(commands, page_number):
        footer_y = 22
        commands.append(
            "0.831 0.686 0.216 RG 0.8 w "
            f"{margin_x} 35 m 806 35 l S"
        )
        _pdf_text(
            commands,
            margin_x,
            footer_y,
            (
                "Dibuat "
                f"{generated_local:%d/%m/%Y %H:%M} · "
                f"Total kerja {_display_minutes(summary['worked_minutes'])} · "
                f"Total lembur {_display_minutes(summary['overtime_minutes'])}"
            ),
            size=7,
            color=(0.38, 0.38, 0.38),
        )
        _pdf_text(
            commands,
            760,
            footer_y,
            f"{page_number} / {total_pages}",
            size=7,
            bold=True,
            color=(0.38, 0.38, 0.38),
        )

    for page_index, page_staff in enumerate(staff_chunks, start=1):
        commands: list[str] = []
        add_page_header(commands, "Ringkasan per Staff")
        table_top = page_height - 116
        x = margin_x
        for title, width in staff_columns:
            _pdf_rect(
                commands,
                x,
                table_top,
                width,
                24,
                fill=(0.831, 0.686, 0.216),
                stroke=(0.72, 0.58, 0.14),
            )
            _pdf_text(
                commands,
                x + 3,
                table_top + 8,
                _truncate_pdf(title, width - 6, font_size=6.2),
                size=6.2,
                bold=True,
            )
            x += width

        if not page_staff:
            _pdf_text(
                commands,
                margin_x,
                table_top - 34,
                "Tidak ada data staff pada periode yang dipilih.",
                size=10,
                color=(0.35, 0.35, 0.35),
            )

        for local_index, staff in enumerate(page_staff):
            global_index = (
                (page_index - 1) * staff_rows_per_page
                + local_index
                + 1
            )
            y = table_top - ((local_index + 1) * staff_row_height)
            row = [
                global_index,
                staff["name"],
                staff["role"],
                staff["days"],
                staff["present"],
                staff["ontime"],
                staff["late"],
                staff["early_leave"],
                staff["overtime"],
                staff["absent"],
                staff["leave"],
                staff["holiday"],
                staff["wfh"],
                _display_minutes(staff["worked_minutes"]),
                _display_minutes(staff["overtime_minutes"]),
            ]
            x = margin_x
            fill = (
                (1, 1, 1)
                if local_index % 2 == 0
                else (0.965, 0.965, 0.965)
            )
            for value, (_, width) in zip(row, staff_columns):
                _pdf_rect(
                    commands,
                    x,
                    y,
                    width,
                    staff_row_height,
                    fill=fill,
                )
                _pdf_text(
                    commands,
                    x + 3,
                    y + 6.5,
                    _truncate_pdf(value, width - 6, font_size=6.2),
                    size=6.2,
                )
                x += width

        add_page_footer(commands, page_index)
        page_streams.append(
            "\n".join(commands).encode("latin-1", errors="replace")
        )

    for page_index, page_records in enumerate(detail_chunks, start=1):
        commands = []
        add_page_header(commands, "Detail Attendance")
        table_top = page_height - 116
        x = margin_x
        for title, width in detail_columns:
            _pdf_rect(
                commands,
                x,
                table_top,
                width,
                24,
                fill=(0.831, 0.686, 0.216),
                stroke=(0.72, 0.58, 0.14),
            )
            _pdf_text(
                commands,
                x + 4,
                table_top + 8,
                title,
                size=7.2,
                bold=True,
            )
            x += width

        if not page_records:
            _pdf_text(
                commands,
                margin_x,
                table_top - 34,
                "Tidak ada data attendance untuk periode yang dipilih.",
                size=10,
                color=(0.35, 0.35, 0.35),
            )

        for local_index, record in enumerate(page_records):
            global_index = (
                (page_index - 1) * detail_rows_per_page
                + local_index
                + 1
            )
            y = table_top - (
                (local_index + 1) * detail_row_height
            )
            profile = record.user
            row = [
                global_index,
                record.date.strftime("%d/%m/%Y"),
                profile.full_name,
                profile.get_role_display(),
                record.get_work_mode_display(),
                _display_time(record.check_in),
                _display_time(record.check_out),
                _display_minutes(record.worked_minutes),
                _display_minutes(record.overtime_minutes),
                record.get_status_display(),
                record.work_policy.name if record.work_policy else "-",
            ]
            x = margin_x
            fill = (
                (1, 1, 1)
                if local_index % 2 == 0
                else (0.965, 0.965, 0.965)
            )
            for value, (_, width) in zip(row, detail_columns):
                _pdf_rect(
                    commands,
                    x,
                    y,
                    width,
                    detail_row_height,
                    fill=fill,
                )
                _pdf_text(
                    commands,
                    x + 4,
                    y + 6,
                    _truncate_pdf(value, width - 8),
                    size=7.2,
                )
                x += width

        add_page_footer(
            commands,
            len(staff_chunks) + page_index,
        )
        page_streams.append(
            "\n".join(commands).encode("latin-1", errors="replace")
        )

    return _build_pdf_document(page_streams, logo_assets=logo_assets)
